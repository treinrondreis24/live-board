import io,json,os,unittest
from types import SimpleNamespace
from unittest.mock import patch
from PIL import Image
from pypdf import PdfReader
import test_cloud,route_maker,route_google,google_connection

class RouteTests(unittest.TestCase):
    def setUp(self):
        self.s=SimpleNamespace(**vars(test_cloud.cloud.server))
        self.start={'lat':47.5,'lon':19.08,'name':'Keleti','address':'Station Budapest Keleti'}
        self.end={'lat':47.499,'lon':19.081,'name':'IntercityHotel Budapest','address':'Baross tér 7-8, Budapest'}
        buf=io.BytesIO();Image.new('RGB',(1100,700),'#dddddd').save(buf,'PNG');self.image=buf.getvalue()
        self.data={'features':[{'geometry':{'type':'LineString','coordinates':[[19.08,47.5],[19.081,47.499]]},'properties':{'distance':203,'time':162,'legs':[{'steps':[{'distance':110,'instruction':{'text':'Loop naar Baross tér.'}},{'distance':93,'instruction':{'text':'Neem het zebrapad.'}}]}]}}]}
    def route(self):
        with patch.object(route_maker,'request',side_effect=[json.dumps(self.data).encode(),self.image]) as call:
            model=route_maker.calculate(self.s,{'start':self.start,'end':self.end})
            self.assertEqual(call.call_args_list[0].args[2]['mode'],'walk')
            self.assertEqual(call.call_args_list[0].args[2]['lang'],'nl')
            self.assertEqual(call.call_args_list[1].kwargs['body']['geojson']['geometry'],self.data['features'][0]['geometry'])
            return model
    def test_render_and_docx_preserve_route(self):
        model=self.route()
        for fmt in ['A4','A5']:
            b={'id':model['id'],'country':'Hongarije','destination':'Budapest','hotel':'IntercityHotel Budapest','format':fmt,'extra':'Eigen aanvullende informatie.'}
            result=route_maker.render(self.s,b)
            self.assertEqual(result['pages'],1)
            self.assertEqual(result['title'],'Route - IntercityHotel Budapest')
            pdf=PdfReader(self.s.DATA/'pdfs'/(result['asset']+'.pdf'));txt=pdf.pages[0].extract_text()
            for text in ['Neem het zebrapad','203 meter','Geoapify','Eigen aanvullende informatie']:self.assertIn(text,txt)
            from docx import Document
            doc=Document(io.BytesIO(route_google.document(self.s,model,result['routeMaker'])))
            self.assertEqual(len(doc.inline_shapes),1)
            self.assertIn('Neem het zebrapad','\n'.join(p.text for row in doc.tables[0].rows for cell in row.cells for p in cell.paragraphs))
        long=route_maker.render(self.s,{**b,'format':'A5','instructions':'\n'.join(f'Stap {n}: loop verder en controleer de straatnaam.' for n in range(70))})
        reader=PdfReader(self.s.DATA/'pdfs'/(long['asset']+'.pdf'))
        self.assertGreater(len(reader.pages),1)
        self.assertIn('Stap 69',''.join(p.extract_text() for p in reader.pages))
    def test_boundaries_and_private_access(self):
        for ident in ['../x','a'*31,None]:
            with self.assertRaises(ValueError):route_maker.stored(self.s,ident,'.png')
        for lat in ['nan',91,'infinity']:
            with self.assertRaises(ValueError):route_maker.point({**self.start,'lat':lat})
        with self.assertRaises(ValueError):route_maker.photo(base64_string:='not a photo')
        for path in ['/api/routes/status','/api/routes/search','/files/route/'+'a'*32]:
            self.assertEqual(test_cloud.request(path,authorized=False)[0],401)
        page=test_cloud.request('/')[2].decode()
        self.assertIn('/seinhuis/boekjesmaker/api/routes/calculate',page)
        self.assertIn('/seinhuis/boekjesmaker/files/route/',page)
        with patch.dict(os.environ,{'GEOAPIFY_API_KEY':'private-key-never-return'}):
            result=test_cloud.request('/api/routes/status')[2]
            self.assertNotIn(b'private-key',result)
            self.assertTrue(json.loads(result)['configured'])
    def test_missing_route_and_google_permission(self):
        with patch.object(route_maker,'request',return_value=b'{"features":[]}'):
            with self.assertRaisesRegex(ValueError,'geen looproute'):route_maker.calculate(self.s,{'start':self.start,'end':self.end})
        with patch.object(google_connection,'status',return_value={'canCreate':False}):
            with self.assertRaisesRegex(ValueError,'toestemming'):route_google.export(self.s,{})

if __name__=='__main__':unittest.main()
