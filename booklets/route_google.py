"""Create new editable Google Docs from routes; never overwrite source documents."""
import io,json,secrets
from urllib.request import Request,urlopen
from urllib.error import HTTPError,URLError
import google_connection as google
import route_maker

def document(s,model,cfg):
    from docx import Document
    from docx.shared import Pt,RGBColor
    doc=Document();section=doc.sections[0];W,H=s.SIZES[cfg['format']]
    section.page_width=Pt(W);section.page_height=Pt(H)
    section.left_margin=section.right_margin=Pt(28);section.top_margin=Pt(28);section.bottom_margin=Pt(44)
    normal=doc.styles['Normal'];normal.font.name='Open Sans';normal.font.size=Pt(10.5);normal.font.color.rgb=RGBColor(0,0,0);normal.paragraph_format.line_spacing=1.6
    title=doc.add_paragraph();r=title.add_run(cfg['title']);r.font.name='Montserrat';r.font.size=Pt(19);r.bold=True
    doc.add_paragraph(f"Vanaf {model['start']['name']} | circa {model['distance']} meter | {model['minutes']} minuten lopen")
    table=doc.add_table(rows=1,cols=2);table.autofit=False
    left,right=table.rows[0].cells;left.width=Pt((W-56)*.55);right.width=Pt((W-56)*.45)
    left.paragraphs[0].add_run().add_picture(str(route_maker.stored(s,model['id'],'.png')),width=Pt((W-56)*.55-12))
    p=left.add_paragraph(route_maker.ATTRIBUTION);p.runs[0].font.size=Pt(7)
    right.paragraphs[0].add_run('Uw looproute').bold=True
    for n,line in enumerate(cfg['instructions'].splitlines()):
        if line.strip():right.add_paragraph(f'{n+1}. {line}')
    right.add_paragraph('Uw hotel').runs[0].bold=True
    right.add_paragraph(model['end']['address'])
    if cfg['extra']:
        doc.add_paragraph('Goed om te weten').runs[0].bold=True;doc.add_paragraph(cfg['extra'])
    if cfg.get('photo'):doc.add_picture(str(route_maker.stored(s,cfg['photo'],'.jpg')),width=Pt(240))
    buf=io.BytesIO();doc.save(buf);return buf.getvalue()

def export(s,b):
    if not google.status(s.DATA)['canCreate']:raise ValueError('Geef eerst extra Google-toestemming via Google koppelen → Routes opslaan in Google Docs. Schakel ook de Google Drive API in voor uw Google-project.')
    item=s.obj(b['id'])
    if item['kind']!='block' or not item['body'].get('routeMaker'):raise ValueError('Sla eerst een Routemaker-bouwsteen op.')
    if item['revision']!=b.get('revision'):raise ValueError('De bouwsteen is gewijzigd. Herlaad eerst.')
    body=item['body'];old=body.get('routeGoogle')
    if old and old.get('asset')==body['asset']:return {'url':old['url']}
    model,cfg=route_maker.configuration(s,{**body['routeMaker'],'title':item['title']})
    raw=document(s,model,cfg);boundary='booklet_'+secrets.token_hex(16)
    payload=(f'--{boundary}\r\nContent-Type: application/json; charset=UTF-8\r\n\r\n'.encode()+json.dumps({'name':item['title'],'mimeType':'application/vnd.google-apps.document'}).encode()+f'\r\n--{boundary}\r\nContent-Type: application/vnd.openxmlformats-officedocument.wordprocessingml.document\r\n\r\n'.encode()+raw+f'\r\n--{boundary}--\r\n'.encode())
    req=Request('https://www.googleapis.com/upload/drive/v3/files?uploadType=multipart&fields=id',data=payload,headers={'Authorization':'Bearer '+google.token(s.DATA),'Content-Type':'multipart/related; boundary='+boundary})
    try:
        with urlopen(req,timeout=55) as response:result=json.loads(response.read(100000))
    except HTTPError as e:
        if e.code==403:raise ValueError('Google weigert het opslaan. Controleer of de Google Drive API aanstaat en koppel Google opnieuw met toestemming om routes op te slaan.') from None
        raise ValueError('Opslaan in Google Docs is niet bevestigd. Controleer Google Drive voordat u opnieuw probeert.') from None
    except (URLError,TimeoutError,OSError):raise ValueError('Google reageerde niet op tijd. Controleer Google Drive voordat u opnieuw probeert.') from None
    url='https://docs.google.com/document/d/'+result['id']+'/edit'
    with s.LOCK,s.db() as c:
        current=c.execute('SELECT revision FROM objects WHERE id=?',(item['id'],)).fetchone()
        if current and current['revision']==item['revision']:
            body['routeGoogle']={'url':url,'asset':body['asset']}
            c.execute('UPDATE objects SET body=?,revision=revision+1 WHERE id=?',(json.dumps(body,ensure_ascii=False),item['id']))
    return {'url':url}
